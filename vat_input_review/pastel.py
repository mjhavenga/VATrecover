from __future__ import annotations

import csv
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path
from typing import Any

from .config import OrgReviewConfig
from .models import TransactionLine, money
from .tax_periods import vat_period_label


PASTEL_COLUMN_ALIASES = {
    "transaction_id": ["transaction_id", "tx id", "entry no", "reference", "document number", "doc no"],
    "line_id": ["line_id", "line no", "line", "entry line"],
    "date": ["date", "transaction date", "document date"],
    "supplier_id": ["supplier_id", "supplier code", "account code", "vendor code"],
    "supplier_name": ["supplier", "supplier name", "vendor", "vendor name", "account name"],
    "supplier_vat_number": ["supplier vat", "vat number", "tax number", "vat reg no"],
    "account_code": ["gl code", "account", "account code", "nominal code", "ledger code"],
    "account_name": ["gl account", "account name", "nominal account", "ledger account"],
    "tax_type": ["tax type", "tax code", "vat code"],
    "description": ["description", "details", "narration"],
    "net_amount": ["net", "net amount", "exclusive", "amount excl", "taxable amount"],
    "vat_amount": ["vat", "vat amount", "tax", "tax amount"],
    "gross_amount": ["gross", "gross amount", "inclusive", "amount incl", "total"],
}


@dataclass(frozen=True)
class PastelImportResult:
    source_path: Path
    lines: list[TransactionLine]


def load_pastel_export(path: str | Path, config: OrgReviewConfig) -> PastelImportResult:
    source_path = Path(path)
    rows = _read_rows(source_path)
    lines = [normalize_pastel_row(row, config, index) for index, row in enumerate(rows, start=1)]
    return PastelImportResult(source_path=source_path, lines=lines)


def normalize_pastel_row(row: dict[str, Any], config: OrgReviewConfig, row_number: int) -> TransactionLine:
    normalized = {_normalize_header(key): value for key, value in row.items()}
    transaction_date = _parse_date(_lookup(normalized, "date"))
    transaction_id = _lookup(normalized, "transaction_id") or f"{Path(str(row.get('_source_file', 'pastel'))).stem}:{row_number}"
    line_id = _lookup(normalized, "line_id") or f"{transaction_id}:{row_number}"
    net = money(_lookup(normalized, "net_amount"))
    vat = money(_lookup(normalized, "vat_amount"))
    gross_value = _lookup(normalized, "gross_amount")
    gross = money(gross_value) if gross_value not in {None, ""} else money(net + vat)

    return TransactionLine(
        tenant_id=config.tenant_id,
        organisation_name=config.organisation_name,
        transaction_type="PASTEL_PURCHASE",
        transaction_id=str(transaction_id),
        line_id=str(line_id),
        document_number=str(_lookup(normalized, "transaction_id") or transaction_id),
        transaction_date=transaction_date,
        vat_period=vat_period_label(transaction_date, config.vat_period_basis),
        supplier_id=_as_optional_str(_lookup(normalized, "supplier_id")),
        supplier_name=_as_optional_str(_lookup(normalized, "supplier_name")),
        supplier_vat_number=_as_optional_str(_lookup(normalized, "supplier_vat_number")),
        account_id=None,
        account_code=_as_optional_str(_lookup(normalized, "account_code")),
        account_name=_as_optional_str(_lookup(normalized, "account_name")),
        tax_type=_as_optional_str(_lookup(normalized, "tax_type")),
        description=_as_optional_str(_lookup(normalized, "description")),
        net_amount=net,
        vat_amount=vat,
        gross_amount=gross,
        source_system="pastel",
        raw={"row": row},
    )


def _read_rows(path: Path) -> list[dict[str, Any]]:
    suffix = path.suffix.lower()
    if suffix == ".csv":
        with path.open(newline="", encoding="utf-8-sig") as handle:
            rows = list(csv.DictReader(handle))
    elif suffix in {".xlsx", ".xlsm"}:
        try:
            from openpyxl import load_workbook
        except ImportError as exc:  # pragma: no cover
            raise RuntimeError("openpyxl is required to import Pastel Excel exports.") from exc
        workbook = load_workbook(path, read_only=True, data_only=True)
        sheet = workbook.active
        values = list(sheet.iter_rows(values_only=True))
        if not values:
            rows = []
        else:
            headers = [str(value or "").strip() for value in values[0]]
            rows = [dict(zip(headers, row)) for row in values[1:] if any(cell is not None for cell in row)]
    else:
        raise ValueError(
            f"Unsupported Pastel file type '{path.suffix}'. Export the Pastel review data to CSV or XLSX first."
        )

    for row in rows:
        row["_source_file"] = str(path)
    return rows


def _lookup(row: dict[str, Any], field_name: str) -> Any:
    for alias in PASTEL_COLUMN_ALIASES[field_name]:
        value = row.get(_normalize_header(alias))
        if value not in {None, ""}:
            return value
    return None


def _normalize_header(value: str) -> str:
    return " ".join(str(value).strip().lower().replace("_", " ").split())


def _parse_date(value: Any) -> date:
    if isinstance(value, date):
        return value
    if value is None:
        raise ValueError("Pastel row is missing a transaction date.")
    text = str(value).strip()
    for date_format in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%Y/%m/%d"):
        try:
            return datetime.strptime(text, date_format).date()
        except ValueError:
            pass
    return datetime.fromisoformat(text).date()


def _as_optional_str(value: Any) -> str | None:
    if value in {None, ""}:
        return None
    return str(value).strip()
