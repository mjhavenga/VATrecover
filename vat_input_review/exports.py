from __future__ import annotations

import csv
import json
from dataclasses import dataclass
from datetime import date
from pathlib import Path

from .io import transaction_line_to_dict
from .models import TransactionLine


EXPORT_COLUMNS = [
    "source_system",
    "tenant_id",
    "organisation_name",
    "transaction_type",
    "transaction_id",
    "line_id",
    "document_number",
    "transaction_date",
    "vat_period",
    "supplier_id",
    "supplier_name",
    "supplier_vat_number",
    "account_id",
    "account_code",
    "account_name",
    "tax_type",
    "description",
    "net_amount",
    "vat_amount",
    "gross_amount",
    "currency_code",
    "source_url",
    "ai_review_prompt",
]


@dataclass(frozen=True)
class TransactionExportPaths:
    csv_path: Path | None = None
    json_path: Path | None = None
    xlsx_path: Path | None = None


def export_transaction_pack(
    lines: list[TransactionLine],
    output_dir: str | Path,
    tenant_id: str,
    run_range: tuple[date, date],
    formats: set[str] | None = None,
) -> TransactionExportPaths:
    formats = formats or {"csv", "json", "xlsx"}
    org_dir = Path(output_dir) / _safe_path_part(tenant_id)
    org_dir.mkdir(parents=True, exist_ok=True)
    date_label = f"{run_range[0].isoformat()}_to_{run_range[1].isoformat()}"
    stem = f"transactions_for_ai_review_{date_label}"

    csv_path = org_dir / f"{stem}.csv" if "csv" in formats else None
    json_path = org_dir / f"{stem}.json" if "json" in formats else None
    xlsx_path = org_dir / f"{stem}.xlsx" if "xlsx" in formats else None

    sorted_lines = sorted(lines, key=lambda line: (line.transaction_date, line.transaction_id, line.line_id))
    if csv_path:
        _write_csv(csv_path, sorted_lines)
    if json_path:
        _write_json(json_path, sorted_lines)
    if xlsx_path:
        _write_xlsx(xlsx_path, sorted_lines, run_range)

    return TransactionExportPaths(csv_path=csv_path, json_path=json_path, xlsx_path=xlsx_path)


def transaction_export_row(line: TransactionLine) -> dict[str, str]:
    row = transaction_line_to_dict(line)
    row["ai_review_prompt"] = (
        "Review this purchase transaction for potential under-claimed input VAT. "
        "Check supplier VAT registration, tax code, account pattern, blocked input VAT, exempt/zero-rated nature, "
        "apportionment risk, and source-document evidence before recommending any correction."
    )
    return {column: str(row.get(column) or "") for column in EXPORT_COLUMNS}


def _write_csv(path: Path, lines: list[TransactionLine]) -> None:
    with path.open("w", newline="", encoding="utf-8-sig") as handle:
        writer = csv.DictWriter(handle, fieldnames=EXPORT_COLUMNS)
        writer.writeheader()
        for line in lines:
            writer.writerow(transaction_export_row(line))


def _write_json(path: Path, lines: list[TransactionLine]) -> None:
    payload = [
        {
            **transaction_line_to_dict(line),
            "ai_review_instruction": (
                "Assess this transaction for input VAT recoverability. Return review notes, evidence checks, "
                "questions for the accountant, and a suggested review priority. Do not mark it as automatically claimable."
            ),
        }
        for line in lines
    ]
    path.write_text(json.dumps(payload, indent=2), encoding="utf-8")


def _write_xlsx(path: Path, lines: list[TransactionLine], run_range: tuple[date, date]) -> None:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError as exc:  # pragma: no cover
        raise RuntimeError("openpyxl is required to export XLSX transaction packs.") from exc

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Transactions"
    sheet["A1"] = "Transactions for AI VAT review"
    sheet["A2"] = f"Period: {run_range[0].isoformat()} to {run_range[1].isoformat()}"
    sheet["A3"] = "Export only. No transactions are changed in the accounting system."
    sheet["A1"].font = Font(bold=True, size=14)
    sheet["A3"].font = Font(italic=True)

    header_row = 5
    for col, header in enumerate(EXPORT_COLUMNS, start=1):
        cell = sheet.cell(row=header_row, column=col, value=header)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")

    for row_idx, line in enumerate(lines, start=header_row + 1):
        row = transaction_export_row(line)
        for col_idx, header in enumerate(EXPORT_COLUMNS, start=1):
            sheet.cell(row=row_idx, column=col_idx, value=row[header])

    widths = {
        "A": 14,
        "B": 38,
        "C": 28,
        "D": 18,
        "E": 38,
        "F": 38,
        "G": 18,
        "H": 14,
        "I": 18,
        "K": 28,
        "L": 18,
        "O": 24,
        "P": 16,
        "Q": 36,
        "W": 72,
    }
    for index in range(1, len(EXPORT_COLUMNS) + 1):
        letter = get_column_letter(index)
        sheet.column_dimensions[letter].width = widths.get(letter, 16)
    sheet.freeze_panes = "A6"
    sheet.auto_filter.ref = f"A{header_row}:W{max(header_row + 1, header_row + len(lines))}"
    workbook.save(path)


def _safe_path_part(value: str) -> str:
    return "".join(ch if ch.isalnum() or ch in {"-", "_"} else "_" for ch in value)
