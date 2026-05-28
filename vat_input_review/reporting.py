from __future__ import annotations

from collections import defaultdict
from dataclasses import dataclass
from datetime import date
from decimal import Decimal
from pathlib import Path

from .config import OrgReviewConfig
from .models import ReviewFlag, money


@dataclass(frozen=True)
class ReportPaths:
    working_paper_xlsx: Path
    client_summary_md: Path


def write_reports(flags: list[ReviewFlag], config: OrgReviewConfig, output_dir: str | Path, run_range: tuple[date, date]) -> ReportPaths:
    org_dir = Path(output_dir) / _safe_path_part(config.tenant_id)
    org_dir.mkdir(parents=True, exist_ok=True)
    date_label = f"{run_range[0].isoformat()}_to_{run_range[1].isoformat()}"
    working_paper = org_dir / f"vat_input_review_working_paper_{date_label}.xlsx"
    summary = org_dir / f"vat_input_review_client_summary_{date_label}.md"
    write_working_paper(flags, config, working_paper, run_range)
    write_client_summary(flags, config, summary, run_range)
    return ReportPaths(working_paper_xlsx=working_paper, client_summary_md=summary)


def write_working_paper(flags: list[ReviewFlag], config: OrgReviewConfig, path: str | Path, run_range: tuple[date, date]) -> None:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError as exc:  # pragma: no cover
        raise RuntimeError("openpyxl is required to write the Excel working paper.") from exc

    sorted_flags = sorted(flags, key=lambda flag: flag.estimated_underclaim, reverse=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "Review Items"

    title = config.organisation_name or config.tenant_id
    ws["A1"] = f"VAT input-tax review: {title}"
    ws["A2"] = f"Review period: {run_range[0].isoformat()} to {run_range[1].isoformat()}"
    ws["A3"] = "All items are review items requiring sign-off before any VAT correction or refund claim."
    ws["A1"].font = Font(bold=True, size=14)
    ws["A3"].font = Font(italic=True)

    headers = [
        "Reviewed Y/N",
        "Tenant ID",
        "Source System",
        "Transaction ID",
        "Line ID",
        "Document No.",
        "Date",
        "VAT Period",
        "Supplier",
        "Supplier VAT No.",
        "Account Code",
        "Account Name",
        "Tax Type",
        "Description",
        "Net",
        "VAT Claimed",
        "VAT Expected",
        "Estimated Under-claim",
        "Reason Code",
        "Confidence",
        "Audit Trail",
    ]
    start_row = 5
    for col, header in enumerate(headers, start=1):
        cell = ws.cell(row=start_row, column=col, value=header)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.alignment = Alignment(wrap_text=True)

    for row_idx, flag in enumerate(sorted_flags, start=start_row + 1):
        line = flag.line
        values = [
            "",
            line.tenant_id,
            line.source_system,
            line.transaction_id,
            line.line_id,
            line.document_number,
            line.transaction_date,
            line.vat_period,
            line.supplier_name,
            line.supplier_vat_number,
            line.account_code,
            line.account_name,
            line.tax_type,
            line.description,
            float(line.net_amount),
            float(line.vat_amount),
            float(flag.expected_vat),
            float(flag.estimated_underclaim),
            flag.reason_code,
            float(flag.confidence),
            _audit_trail(flag),
        ]
        for col, value in enumerate(values, start=1):
            ws.cell(row=row_idx, column=col, value=value)

    total_row = start_row + len(sorted_flags) + 1
    ws.cell(row=total_row, column=17, value="Total potential recovery").font = Font(bold=True)
    ws.cell(row=total_row, column=18, value=f"=SUM(R{start_row + 1}:R{total_row - 1})").font = Font(bold=True)

    money_columns = [15, 16, 17, 18]
    for col in money_columns:
        for row in range(start_row + 1, total_row + 1):
            ws.cell(row=row, column=col).number_format = '#,##0.00'
    for row in range(start_row + 1, total_row):
        ws.cell(row=row, column=20).number_format = "0%"

    widths = {
        1: 14,
        3: 14,
        4: 38,
        5: 38,
        6: 18,
        7: 12,
        8: 18,
        9: 28,
        10: 18,
        11: 14,
        12: 24,
        13: 16,
        14: 36,
        18: 20,
        19: 34,
        21: 60,
    }
    for col in range(1, len(headers) + 1):
        ws.column_dimensions[get_column_letter(col)].width = widths.get(col, 16)
    ws.freeze_panes = "A6"
    ws.auto_filter.ref = f"A{start_row}:U{max(start_row + 1, total_row - 1)}"

    summary = wb.create_sheet("Summary")
    _write_summary_sheet(summary, sorted_flags, config, run_range)
    wb.save(path)


def write_client_summary(flags: list[ReviewFlag], config: OrgReviewConfig, path: str | Path, run_range: tuple[date, date]) -> None:
    total = sum((flag.estimated_underclaim for flag in flags), money("0"))
    by_reason = _breakdown(flags, lambda flag: flag.reason_code)
    by_period = _breakdown(flags, lambda flag: flag.line.vat_period)
    org_name = config.organisation_name or config.tenant_id

    lines = [
        f"# VAT Input-Tax Review Summary: {org_name}",
        "",
        f"Review period: {run_range[0].isoformat()} to {run_range[1].isoformat()}",
        "",
        "This report identifies potential under-claimed input VAT review items only. Each item must be checked against the source tax invoice and signed off before any correction or refund claim is submitted.",
        "",
        f"Total potential recovery: ZAR {total:,.2f}",
        f"Review item count: {len(flags)}",
        "",
        "## Breakdown by reason",
        "",
        "| Reason code | Items | Potential recovery |",
        "| --- | ---: | ---: |",
    ]
    lines.extend(f"| {reason} | {count} | ZAR {amount:,.2f} |" for reason, (count, amount) in by_reason.items())
    lines.extend(["", "## Breakdown by VAT period", "", "| VAT period | Items | Potential recovery |", "| --- | ---: | ---: |"])
    lines.extend(f"| {period} | {count} | ZAR {amount:,.2f} |" for period, (count, amount) in by_period.items())
    lines.extend(
        [
            "",
            "## Reviewer note",
            "",
            "The Excel working paper includes the source document ID, line ID, reason code, confidence, estimated VAT expected, estimated under-claim, and a reviewed Y/N column for sign-off.",
        ]
    )
    Path(path).write_text("\n".join(lines), encoding="utf-8")


def _write_summary_sheet(ws, flags: list[ReviewFlag], config: OrgReviewConfig, run_range: tuple[date, date]) -> None:
    ws["A1"] = "Client summary"
    ws["A1"].font = ws["A1"].font.copy(bold=True, size=14)
    ws["A2"] = config.organisation_name or config.tenant_id
    ws["A3"] = f"{run_range[0].isoformat()} to {run_range[1].isoformat()}"
    ws["A5"] = "Total potential recovery"
    ws["B5"] = float(sum((flag.estimated_underclaim for flag in flags), money("0")))
    ws["A6"] = "Review item count"
    ws["B6"] = len(flags)

    row = 8
    ws.cell(row=row, column=1, value="Reason code").font = ws.cell(row=row, column=1).font.copy(bold=True)
    ws.cell(row=row, column=2, value="Items").font = ws.cell(row=row, column=2).font.copy(bold=True)
    ws.cell(row=row, column=3, value="Potential recovery").font = ws.cell(row=row, column=3).font.copy(bold=True)
    for reason, (count, amount) in _breakdown(flags, lambda flag: flag.reason_code).items():
        row += 1
        ws.cell(row=row, column=1, value=reason)
        ws.cell(row=row, column=2, value=count)
        ws.cell(row=row, column=3, value=float(amount))

    row += 3
    ws.cell(row=row, column=1, value="VAT period").font = ws.cell(row=row, column=1).font.copy(bold=True)
    ws.cell(row=row, column=2, value="Items").font = ws.cell(row=row, column=2).font.copy(bold=True)
    ws.cell(row=row, column=3, value="Potential recovery").font = ws.cell(row=row, column=3).font.copy(bold=True)
    for period, (count, amount) in _breakdown(flags, lambda flag: flag.line.vat_period).items():
        row += 1
        ws.cell(row=row, column=1, value=period)
        ws.cell(row=row, column=2, value=count)
        ws.cell(row=row, column=3, value=float(amount))

    ws.column_dimensions["A"].width = 42
    ws.column_dimensions["B"].width = 16
    ws.column_dimensions["C"].width = 22
    for row_idx in range(5, row + 1):
        ws.cell(row=row_idx, column=3).number_format = '#,##0.00'


def _breakdown(flags: list[ReviewFlag], key_fn) -> dict[str, tuple[int, Decimal]]:
    grouped: dict[str, list[ReviewFlag]] = defaultdict(list)
    for flag in flags:
        grouped[str(key_fn(flag))].append(flag)
    return {
        key: (len(values), sum((flag.estimated_underclaim for flag in values), money("0")))
        for key, values in sorted(grouped.items())
    }


def _audit_trail(flag: ReviewFlag) -> str:
    line = flag.line
    return (
        f"transaction_type={line.transaction_type}; transaction_id={line.transaction_id}; "
        f"line_id={line.line_id}; document_no={line.document_number}; reason={flag.reason}"
    )


def _safe_path_part(value: str) -> str:
    return "".join(ch if ch.isalnum() or ch in {"-", "_"} else "_" for ch in value)
